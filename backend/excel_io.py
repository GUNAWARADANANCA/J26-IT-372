"""Import and export the research workbook layout.

One sheet. Employee columns, then one daily KPI row.
Employee_ID is required. A Date value makes the row a KPI log as well.
"""

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from pymongo.errors import DuplicateKeyError

from db import get_db

EMPLOYEE_COLUMNS = [
    "Employee_ID",
    "Employee_Name",
    "Post_Office",
    "Designation",
    "Work_Category",
    "Years_of_Experience",
    "Employment_Type",
    "Reports_To",
    "Prior_6Month_Avg_KPI_Score",
    "Promotion_Eligible",
    "Last_Promotion_Date",
    "Months_Since_Last_Promotion",
    "Reward_Received",
    "Reward_Type",
]

KPI_COLUMNS = [
    "Employee_ID",
    "Date",
    "Day_of_Week",
    "Route_ID",
    "Route_Length_km",
    "Distance_Covered_km",
    "Primary_Task_Type",
    "Assigned_Work_Count",
    "Completed_Work_Count",
    "Uncompleted_Work_Count",
    "Completion_Rate",
    "Correct_Work_Count",
    "Incorrect_Work_Count",
    "Accuracy_Rate",
    "On_Time_Work_Count",
    "Late_Work_Count",
    "Timeliness_Rate",
    "Approvals_Processed",
    "Accounts_Verified",
    "Revenue_Collected_LKR",
    "KPI_Score",
    "Performance_Category",
]

# Same column order as the imported research workbook.
WORKBOOK_COLUMNS = EMPLOYEE_COLUMNS + KPI_COLUMNS[1:]

_HEADER_INDEX = {name.lower(): name for name in WORKBOOK_COLUMNS}


def _canonical_header(value) -> str | None:
    if value is None:
        return None
    key = str(value).strip().replace(" ", "_").lower()
    return _HEADER_INDEX.get(key)


def _cell(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value == int(value) and abs(value) < 1e15:
        return int(value)
    if isinstance(value, str):
        text = value.strip()
        return text or None
    return value


def _text_id(value) -> str:
    parsed = _cell(value)
    if parsed is None:
        return ""
    return str(parsed).strip()


def _excel_value(value):
    if value is None or value == "":
        return None
    if isinstance(value, str) and len(value) >= 10 and value[4:5] == "-" and value[7:8] == "-":
        try:
            return date.fromisoformat(value[:10])
        except ValueError:
            return value
    return value


def _store(collection, doc: dict) -> None:
    try:
        collection.replace_one({"_id": doc["_id"]}, doc, upsert=True)
    except DuplicateKeyError:
        collection.replace_one({"_id": doc["_id"]}, doc, upsert=True)


def import_workbook(payload: bytes) -> dict:
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as exc:
        raise ValueError("The workbook is empty") from exc

    columns = [_canonical_header(cell) for cell in header_row]
    if "Employee_ID" not in columns:
        raise ValueError("The first row must include an Employee_ID column")

    has_date = "Date" in columns
    employees = get_db()["employee"]
    logs = get_db()["daily_kpi_log"]
    employee_count = 0
    kpi_count = 0
    skipped = 0

    for raw in rows:
        record = {}
        for index, column in enumerate(columns):
            if column is None:
                continue
            record[column] = _cell(raw[index] if index < len(raw) else None)

        employee_id = _text_id(record.get("Employee_ID"))
        if not employee_id:
            skipped += 1
            continue

        current = employees.find_one({"_id": employee_id}) or {"_id": employee_id}
        current["name"] = employee_id
        current["employee_id"] = employee_id
        for column in EMPLOYEE_COLUMNS:
            if column not in columns or column == "Employee_ID":
                continue
            current[column.lower()] = record.get(column)
        _store(employees, current)
        employee_count += 1

        if not has_date or record.get("Date") in (None, ""):
            continue

        log_date = str(record["Date"])[:10]
        log_name = f"{employee_id}-{log_date}"
        existing = logs.find_one({"_id": log_name}) or {"_id": log_name}
        existing["name"] = log_name
        existing["employee_id"] = employee_id
        existing["date"] = log_date
        for column in KPI_COLUMNS:
            if column not in columns or column in ("Employee_ID", "Date"):
                continue
            existing[column.lower()] = record.get(column)
        _store(logs, existing)
        kpi_count += 1

    workbook.close()
    return {
        "employees_upserted": employee_count,
        "kpi_logs_upserted": kpi_count,
        "skipped": skipped,
    }


def _sheet(title: str, columns: list[str], docs: list[dict]):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(columns)
    for doc in docs:
        sheet.append([_excel_value(doc.get(column.lower())) for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_employees() -> bytes:
    docs = list(get_db()["employee"].find().sort("employee_id", 1))
    return _sheet("Employees", EMPLOYEE_COLUMNS, docs)


def export_kpi_logs() -> bytes:
    employees = {
        doc["_id"]: doc for doc in get_db()["employee"].find()
    }
    logs = list(get_db()["daily_kpi_log"].find().sort([("employee_id", 1), ("date", 1)]))
    rows = []
    for log in logs:
        employee = employees.get(log.get("employee_id"), {})
        merged = {**employee, **log}
        merged["employee_id"] = log.get("employee_id")
        rows.append(merged)
    return _sheet("Daily KPI", WORKBOOK_COLUMNS, rows)
